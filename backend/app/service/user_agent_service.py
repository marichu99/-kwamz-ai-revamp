import logging
from app import db
from app.model.useragent import UserAgent, user_agent_companies
from app.model.agentcompany import AgentCompany
from sqlalchemy import exists, or_
from app.model.company import Company
from app.utils.gcs_storage import upload_agent_image, delete_agent_image
from datetime import datetime
from openpyxl import load_workbook
from io import BytesIO
import os
import re
from flask import current_app
from sqlalchemy.exc import IntegrityError

logger = logging.getLogger(__name__)

class UserAgentService:
    def __init__(self):
        self.allowed_extensions = {'jpg', 'jpeg', 'png', 'gif'}
        self.allowed_extensions_xls = {'xlsx', 'xls'}
        self.max_file_size = 5 * 1024 * 1024  # 5MB

    def allowed_file(self, filename):
        """Check if the file extension is allowed for images."""
        return '.' in filename and filename.rsplit('.', 1)[1].lower() in self.allowed_extensions

    def allowed_file_xls(self, filename):
        """Check if the file extension is allowed for Excel files."""
        return '.' in filename and filename.rsplit('.', 1)[1].lower() in self.allowed_extensions_xls

    def get_upload_dir(self):
        """Get the upload directory for profile images."""
        upload_dir = os.path.join(current_app.root_path, 'useragents', 'profiles')
        os.makedirs(upload_dir, exist_ok=True)
        return upload_dir

    def _serialize_primary_company(self, primary):
        if not primary:
            return None
        top_org_name = None
        if primary.top_organization:
            parent = Company.query.filter_by(shortcode=primary.top_organization).first()
            if parent:
                top_org_name = parent.company_name
        return {
            'id': primary.id,
            'company_name': primary.company_name,
            'short_code': primary.short_code,
            'is_scraped': primary.last_scraped_at is not None,
            'last_scraped_at': primary.last_scraped_at.isoformat() if primary.last_scraped_at else None,
            'identity_status': primary.identity_status,
            'top_organization': primary.top_organization,
            'top_organization_name': top_org_name,
            'parent_short_code': primary.parent_short_code,
        }

    def serialize_user_agent(self, user_agent, requesting_user_id=None):
        """Serialize UserAgent object to JSON.

        requesting_user_id: when provided, agent_companies is filtered to only
        include AgentCompanies that belong to that user, so operators shared
        across multiple onboarded companies never expose another user's tills.
        """
        primary = user_agent.agent_company

        # JWT identities arrive as strings — normalize before comparing to int columns
        if requesting_user_id is not None:
            requesting_user_id = int(requesting_user_id)

        # Scope M2M assignments to the requesting user's own companies only
        all_companies = user_agent.agent_companies
        if requesting_user_id is not None:
            all_companies = [
                ac for ac in all_companies
                if ac.user_id == requesting_user_id
                or (ac.company and ac.company.user_id == requesting_user_id)
            ]

            # Never present another tenant's till as this agent's primary —
            # fall back to the first of the requesting user's own tills.
            if primary is not None:
                primary_is_own = (
                    primary.user_id == requesting_user_id
                    or (primary.company and primary.company.user_id == requesting_user_id)
                )
                if not primary_is_own:
                    primary = all_companies[0] if all_companies else None

        return {
            'id': user_agent.id,
            'firstname': user_agent.firstname,
            'lastname': user_agent.lastname,
            'idnumber': user_agent.idnumber,
            'phone_number': user_agent.phone_number,
            'is_authentic': user_agent.is_authentic,
            'authenticity_desc': user_agent.authenticity_desc,
            'image_loc': user_agent.image_loc,
            'date_of_birth': user_agent.date_of_birth.isoformat() if user_agent.date_of_birth else None,
            'operator_role': user_agent.operator_role,
            'agent_company_id': primary.id if primary else None,
            'agent_company_name': primary.company_name if primary else None,
            'agent_company_ids': [ac.id for ac in all_companies],
            'agent_company_names': [ac.company_name for ac in all_companies],
            'agent_companies': [{
                'id': company.id,
                'company_name': company.company_name,
                'registration_number': company.registration_number,
                'short_code': company.short_code,
                'is_scraped': company.last_scraped_at is not None,
                'last_scraped_at': company.last_scraped_at.isoformat() if company.last_scraped_at else None,
                'identity_status': company.identity_status,
                'top_organization': company.top_organization,
                'parent_short_code': company.parent_short_code,
                'is_primary': company.id == (primary.id if primary else None)
            } for company in all_companies],
            'primary_company': self._serialize_primary_company(primary),
        }

    def create_user_agent(self, data, files, user_id):
        """Create a new UserAgent with optional file upload."""
        try:
            # Validate required fields
            required_fields = ['firstname', 'lastname', 'idnumber']
            missing_fields = [field for field in required_fields if not data.get(field)]
            if missing_fields:
                return None, f"Missing required fields: {', '.join(missing_fields)}"

            # Parse primary agent_company_id
            primary_company_id = data.get('agent_company_id')
            if primary_company_id:
                primary_company_id = int(primary_company_id)

            # Parse additional agent company IDs (multi-select, optional)
            company_ids = []
            agent_company_ids = data.get('agent_company_ids[]') or data.get('agent_company_ids')
            if agent_company_ids:
                if isinstance(agent_company_ids, str):
                    if ',' in agent_company_ids:
                        company_ids = [int(id.strip()) for id in agent_company_ids.split(',') if id.strip()]
                    elif agent_company_ids.strip():
                        company_ids = [int(agent_company_ids)]
                else:
                    company_ids = [int(id) for id in agent_company_ids]

            # Require at least a primary company or additional companies
            if not primary_company_id and not company_ids:
                return None, 'At least one agent company is required'

            # Fall back: if no primary set, use first from additional list
            if not primary_company_id and company_ids:
                primary_company_id = company_ids[0]

            # Ensure primary company is included in the additional list for m2m
            if primary_company_id and primary_company_id not in company_ids:
                company_ids.append(primary_company_id)

            # Fetch agent companies for m2m relationship
            agent_companies = []
            if company_ids:
                agent_companies = AgentCompany.query.filter(AgentCompany.id.in_(company_ids)).all()
                if len(agent_companies) != len(company_ids):
                    return None, 'One or more agent companies not found'

            # Validate primary company exists
            if primary_company_id:
                primary_company = AgentCompany.query.get(primary_company_id)
                if not primary_company:
                    return None, 'Primary agent company not found'

            # Parse date_of_birth if provided
            date_of_birth = None
            if data.get('date_of_birth'):
                try:
                    date_of_birth = datetime.strptime(data['date_of_birth'], "%Y-%m-%d").date()
                except ValueError:
                    return None, 'Invalid date_of_birth format. Use YYYY-MM-DD'

            # Handle file upload to GCS
            image_loc = None
            if 'image' in files:
                file = files['image']
                if file.filename == '':
                    return None, 'No file selected'
                result = upload_agent_image(file, user_id)
                if not result['success']:
                    return None, result['error']
                image_loc = result['gcs_path']

            # Create new user
            new_user = UserAgent(
                firstname=data['firstname'],
                lastname=data['lastname'],
                idnumber=data['idnumber'],
                phone_number=data.get('phone_number'),
                is_authentic=data.get('is_authentic') == 'true',
                authenticity_desc=data.get('authenticity_desc'),
                image_loc=image_loc,
                date_of_birth=date_of_birth,
                operator_role=data.get('operator_role') or None,
                user_id=user_id,
                agent_company_id=primary_company_id
            )

            # Add agent companies to m2m
            for company in agent_companies:
                new_user.agent_companies.append(company)

            db.session.add(new_user)
            db.session.commit()
            return self.serialize_user_agent(new_user, requesting_user_id=user_id), None

        except ValueError:
            db.session.rollback()
            return None, 'Invalid agent company ID format'
        except IntegrityError as e:
            db.session.rollback()
            return None, 'Duplicate entry detected. ID number, phone, or other unique fields must be unique.'
        except Exception as e:
            db.session.rollback()
            current_app.logger.error(f"Error creating user: {str(e)}")
            return None, 'Internal server error'

    def get_all_users(self):
        """Retrieve all UserAgents."""
        try:
            users = UserAgent.query.all()
            return [self.serialize_user_agent(user) for user in users], None
        except Exception as e:
            current_app.logger.error(f"Error fetching users: {str(e)}")
            return [], 'Internal server error'
        
        
    def get_all_users_by_userid(self, user_id):
        """
        Retrieve active UserAgents for this user.

        Scraped agents (idnumber starts with 'MPESA-') are only shown when
        they still have an active M2M entry in user_agent_companies — meaning
        they haven't been swapped out. Manually-created agents (real ID numbers)
        are always shown.
        """
        try:
            logger.info(f"The user id is {user_id}")
            user_id = int(user_id)
            is_scraped   = UserAgent.idnumber.like('MPESA-%')
            still_linked = exists().where(
                user_agent_companies.c.user_agent_id == UserAgent.id
            )
            # An operator can serve tills for several tenants but the record is
            # unique per ID number, so also include agents owned by another user
            # that are M2M-linked to one of this user's tills.
            on_own_till = UserAgent.agent_companies.any(AgentCompany.user_id == user_id)
            users = (
                UserAgent.query
                .filter(or_(UserAgent.user_id == user_id, on_own_till))
                .filter(
                    or_(
                        ~is_scraped,   # manual agent — always include
                        still_linked,  # scraped agent — only if still on a till
                    )
                )
                .all()
            )
            return [self.serialize_user_agent(user, requesting_user_id=user_id) for user in users], None
        except Exception as e:
            logger.error(f"Error fetching users: {str(e)}")
            current_app.logger.error(f"Error fetching users: {str(e)}")
            return [], 'Internal server error'

    def get_user(self, user_agent_id, requesting_user_id=None):
        """Retrieve a single UserAgent by ID."""
        try:
            user = UserAgent.query.get_or_404(user_agent_id)
            return self.serialize_user_agent(user, requesting_user_id=requesting_user_id), None
        except Exception as e:
            return None, str(e)

    def get_scrape_status(self, user_agent_id, requesting_user_id=None):
        """Check scrape status for all agent companies linked to a user agent."""
        try:
            user_agent = UserAgent.query.get_or_404(user_agent_id)
            all_companies = user_agent.agent_companies
            # Scope to requesting user's own companies only
            if requesting_user_id is not None:
                all_companies = [
                    ac for ac in all_companies
                    if ac.user_id == requesting_user_id
                    or (ac.company and ac.company.user_id == requesting_user_id)
                ]
            companies = all_companies or ([user_agent.agent_company] if user_agent.agent_company else [])

            result = []
            all_scraped = True
            for company in companies:
                is_scraped = company.last_scraped_at is not None
                if not is_scraped:
                    all_scraped = False
                result.append({
                    'id': company.id,
                    'company_name': company.company_name,
                    'short_code': company.short_code,
                    'business_short_code': company.business_short_code,
                    'is_scraped': is_scraped,
                    'last_scraped_at': company.last_scraped_at.isoformat() if company.last_scraped_at else None,
                    'identity_status': company.identity_status,
                    'is_primary': company.id == user_agent.agent_company_id,
                })

            return {
                'user_agent_id': user_agent_id,
                'all_scraped': all_scraped,
                'companies': result,
                'total_companies': len(result),
                'scraped_count': sum(1 for c in result if c['is_scraped']),
                'unscraped_count': sum(1 for c in result if not c['is_scraped']),
            }, None
        except Exception as e:
            current_app.logger.error(f"Error checking scrape status: {str(e)}")
            return None, str(e)

    def update_user(self, user_id, data, files, requesting_user_id=None):
        """Update an existing UserAgent."""
        try:
            user = UserAgent.query.get_or_404(user_id)

            if not data and not files.get('image'):
                return None, 'No data provided'

            # Update fields if provided
            if 'firstname' in data:
                firstname = data['firstname'].strip()
                if not firstname:
                    return None, 'First name cannot be empty'
                user.firstname = firstname

            if 'lastname' in data:
                lastname = data['lastname'].strip()
                if not lastname:
                    return None, 'Last name cannot be empty'
                user.lastname = lastname

            if 'idnumber' in data:
                idnumber = data['idnumber'].strip()
                if not idnumber:
                    return None, 'ID number cannot be empty'
                user.idnumber = idnumber

            if 'phone_number' in data:
                user.phone_number = data['phone_number'] or None

            if 'is_authentic' in data:
                user.is_authentic = data['is_authentic'] == 'true'

            if 'authenticity_desc' in data:
                user.authenticity_desc = data['authenticity_desc'] or None

            if 'date_of_birth' in data:
                date_of_birth = data['date_of_birth']
                if date_of_birth:
                    try:
                        user.date_of_birth = datetime.strptime(date_of_birth, "%Y-%m-%d").date()
                    except ValueError:
                        return None, 'Invalid date_of_birth format. Use YYYY-MM-DD'
                else:
                    user.date_of_birth = None

            if 'operator_role' in data:
                user.operator_role = data['operator_role'] or None

            # Handle primary agent company update
            if 'agent_company_id' in data:
                agent_company_id = data.get('agent_company_id')
                if agent_company_id:
                    user.agent_company_id = int(agent_company_id)
                else:
                    user.agent_company_id = None

            # Handle agent companies update (m2m)
            if 'agent_company_ids[]' in data or 'agent_company_ids' in data:
                agent_company_ids = data.get('agent_company_ids[]') or data.get('agent_company_ids')
                company_ids = []
                if agent_company_ids:
                    if isinstance(agent_company_ids, str):
                        if ',' in agent_company_ids:
                            company_ids = [int(id.strip()) for id in agent_company_ids.split(',') if id.strip()]
                        elif agent_company_ids.strip():
                            company_ids = [int(agent_company_ids)]
                    else:
                        company_ids = [int(id) for id in agent_company_ids]

                # Ensure primary company is in the m2m list
                if user.agent_company_id and user.agent_company_id not in company_ids:
                    company_ids.append(user.agent_company_id)

                if company_ids:
                    agent_companies = AgentCompany.query.filter(AgentCompany.id.in_(company_ids)).all()
                    if len(agent_companies) != len(company_ids):
                        return None, 'One or more agent companies not found'
                    user.agent_companies.clear()
                    for company in agent_companies:
                        user.agent_companies.append(company)
                else:
                    user.agent_companies.clear()

            # Handle file upload to GCS
            if 'image' in files:
                file = files['image']
                if file.filename == '':
                    return None, 'No file selected'
                # Delete old image from GCS if present
                if user.image_loc:
                    delete_agent_image(user.image_loc)
                result = upload_agent_image(file, user_id)
                if not result['success']:
                    return None, result['error']
                user.image_loc = result['gcs_path']

            db.session.commit()
            return self.serialize_user_agent(user, requesting_user_id=requesting_user_id), None

        except ValueError:
            db.session.rollback()
            return None, 'Invalid agent company ID format'
        except IntegrityError as e:
            db.session.rollback()
            error_msg = str(e.orig).lower()
            if 'unique constraint' in error_msg or 'duplicate key' in error_msg:
                return None, 'Duplicate ID number or other unique field'
            return None, 'Database error occurred'
        except Exception as e:
            db.session.rollback()
            current_app.logger.error(f"Error updating user {user_id}: {str(e)}")
            return None, 'Internal server error'

    def delete_user(self, user_id):
        """Delete a UserAgent."""
        try:
            user = UserAgent.query.get_or_404(user_id)
            if user.image_loc:
                delete_agent_image(user.image_loc)
            db.session.delete(user)
            db.session.commit()
            return True, None
        except Exception as e:
            db.session.rollback()
            current_app.logger.error(f"Error deleting user: {str(e)}")
            return False, 'Internal server error'

    def validate_batch(self, file):
        """Validate an Excel file for batch user creation."""
        try:
            if not file or file.filename == '':
                return None, None, 'No file selected'
            if not self.allowed_file_xls(file.filename):
                return None, None, 'Invalid file type. Only .xlsx and .xls files are allowed'

            file_stream = BytesIO(file.read())
            workbook = load_workbook(file_stream)
            sheet = workbook.active

            headers = [cell.value.strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1)) if cell.value]
            data_rows = [
                [cell.value for cell in row]
                for row in sheet.iter_rows(min_row=2)
                if any(cell.value for cell in row)
            ]

            errors = []
            valid_users = []
            seen_idnumbers = set()
            seen_phones = set([u.phone_number for u in UserAgent.query.filter(UserAgent.phone_number.isnot(None)).all()])

            for row_index, row in enumerate(data_rows, start=2):
                if len(row) < 5:
                    errors.append({
                        'rowIndex': row_index,
                        'firstname': row[0] if len(row) > 0 else None,
                        'lastname': row[1] if len(row) > 1 else None,
                        'idnumber': row[2] if len(row) > 2 else None,
                        'phone_number': row[3] if len(row) > 3 else None,
                        'store_number': row[4] if len(row) > 4 else None,
                        'reason': 'Incomplete row (requires firstname, lastname, idnumber, phone_number, store_number)'
                    })
                    continue

                user_data = {
                    'firstname': str(row[0] or '').strip(),
                    'lastname': str(row[1] or '').strip(),
                    'idnumber': str(row[2] or '').strip(),
                    'phone_number': str(row[3] or '').strip() if row[3] else None,
                    'store_number': str(row[4] or '').strip(),
                    'rowIndex': row_index
                }

                if not user_data['firstname'] or not user_data['lastname'] or not user_data['idnumber'] or not user_data['store_number']:
                    errors.append({**user_data, 'reason': 'Missing required fields (firstname, lastname, idnumber, store_number)'})
                    continue

                if not re.match(r'^[A-Za-z\s-]+$', user_data['firstname']) or not re.match(r'^[A-Za-z\s-]+$', user_data['lastname']):
                    errors.append({**user_data, 'reason': 'Invalid name format (only letters, spaces, and hyphens allowed)'})
                    continue

                if user_data['phone_number']:
                    if not re.match(r'^\+?\d{1,3}?\d{9,12}$', user_data['phone_number']):
                        errors.append({**user_data, 'reason': 'Invalid phone number format (must be 9-12 digits, optional + and 1-3 digit country code)'})
                        continue

                if user_data['idnumber'] in seen_idnumbers:
                    errors.append({**user_data, 'reason': 'Duplicate ID number within file'})
                    continue
                seen_idnumbers.add(user_data['idnumber'])

                if UserAgent.query.filter_by(idnumber=user_data['idnumber']).first():
                    errors.append({**user_data, 'reason': 'ID number already exists in database'})
                    continue

                if user_data['phone_number'] and user_data['phone_number'] in seen_phones:
                    errors.append({**user_data, 'reason': 'Phone number already exists in database'})
                    continue
                seen_phones.add(user_data['phone_number'])

                company = AgentCompany.query.filter_by(store_number=user_data['store_number']).first()
                if not company:
                    errors.append({**user_data, 'reason': f"Invalid Store number '{user_data['store_number']}' (no matching company)"})
                    continue

                user_data['company_id'] = company.id
                user_data['company_name'] = company.company_name

                valid_users.append(user_data)

            return valid_users, errors, None

        except Exception as e:
            current_app.logger.error(f"Error validating batch: {str(e)}")
            return None, None, f'Failed to validate batch: {str(e)}'

    def batch_create_useragents(self, file):
        """Create UserAgents from an Excel file."""
        try:
            if not file or file.filename == '':
                return None, 0, 'No file selected'
            if not self.allowed_file_xls(file.filename):
                return None, 0, 'Invalid file type. Only .xlsx files are allowed'

            file_stream = BytesIO(file.read())
            workbook = load_workbook(file_stream)
            sheet = workbook.active

            headers = [cell.value.strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1)) if cell.value]
            data_rows = [
                [cell.value for cell in row]
                for row in sheet.iter_rows(min_row=2)
                if any(cell.value for cell in row)
            ]

            created_users = []
            errors = []

            for row in data_rows:
                if len(row) < 5:
                    errors.append(f"Invalid row (too few columns): {row}")
                    continue

                store_number = str(row[4] or '') if row[4] else None

                user_data = {
                    'firstname': str(row[0] or ''),
                    'lastname': str(row[1] or ''),
                    'idnumber': str(row[2] or ''),
                    'phone_number': str(row[3] or '') if row[3] else None,
                    'is_authentic': False
                }

                if not user_data['firstname'] or not user_data['lastname'] or not user_data['idnumber']:
                    errors.append(f"Missing required fields in row: {user_data}")
                    continue

                if UserAgent.query.filter_by(idnumber=user_data['idnumber']).first():
                    errors.append(f"Duplicate ID number: {user_data['idnumber']}")
                    continue

                try:
                    new_user = UserAgent(**user_data)
                    existing_company_store = AgentCompany.query.filter_by(store_number=store_number).first()
                    if existing_company_store:
                        new_user.agent_company_id = existing_company_store.id
                        new_user.agent_companies.append(existing_company_store)
                    db.session.add(new_user)
                    created_users.append(user_data)
                except Exception as e:
                    errors.append(f"Error creating user {user_data['idnumber']}: {str(e)}")
                    continue

            try:
                db.session.commit()
                return created_users, len(created_users), None
            except Exception as e:
                db.session.rollback()
                return None, 0, f'Failed to save users: {str(e)}'

        except Exception as e:
            db.session.rollback()
            return None, 0, f'Failed to process batch upload: {str(e)}'